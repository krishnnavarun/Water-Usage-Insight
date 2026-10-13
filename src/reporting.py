from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / 'outputs'


def build_insights(df: pd.DataFrame) -> dict:
    df = df.copy()
    locality = df.groupby('locality')['usage'].mean().sort_values(ascending=False)
    connection_type = df.groupby('connection_type')['usage'].mean().sort_values(ascending=False)
    repair = df[df['leak_report'] == 1].groupby('connection_id')['usage'].agg(['mean', 'count'])
    anomaly_rows = df[df['model_anomaly'] == 1].sort_values('anomaly_score', ascending=False)
    return {
        'row_count': int(len(df)),
        'date_start': str(df['date'].min().date()),
        'date_end': str(df['date'].max().date()),
        'mean_usage': float(df['usage'].mean()),
        'median_usage': float(df['usage'].median()),
        'top_locality': str(locality.index[0]),
        'top_locality_mean_usage': float(locality.iloc[0]),
        'connection_type_means': {str(k): float(v) for k, v in connection_type.items()},
        'reported_leak_readings': int(df['leak_report'].sum()),
        'model_anomalies': int(df['model_anomaly'].sum()),
        'top_anomaly_connection': str(anomaly_rows.iloc[0]['connection_id']) if not anomaly_rows.empty else None,
        'repair_connection_count': int(len(repair)),
    }


def create_charts(df: pd.DataFrame) -> None:
    sns.set_theme(style='whitegrid')
    chart_dir = OUTPUT_DIR / 'charts'
    chart_dir.mkdir(parents=True, exist_ok=True)
    daily = df.groupby('date', as_index=False)['usage'].mean()
    plt.figure(figsize=(10, 4))
    plt.plot(daily['date'], daily['usage'], color='#167c80', linewidth=1)
    plt.title('Average Daily Water Usage')
    plt.xlabel('Date')
    plt.ylabel('Usage')
    plt.tight_layout()
    plt.savefig(chart_dir / 'daily_usage.png', dpi=150)
    plt.close()

    plt.figure(figsize=(8, 4))
    locality = df.groupby('locality', as_index=False)['usage'].mean().sort_values('usage', ascending=False)
    sns.barplot(data=locality, x='usage', y='locality', color='#35a7a0')
    plt.title('Average Usage by Locality')
    plt.xlabel('Average usage')
    plt.ylabel('')
    plt.tight_layout()
    plt.savefig(chart_dir / 'locality_usage.png', dpi=150)
    plt.close()

    plt.figure(figsize=(8, 4))
    sns.scatterplot(data=df.sample(min(5000, len(df)), random_state=42), x='temperature', y='usage', hue='connection_type', alpha=0.45)
    plt.title('Temperature and Water Usage')
    plt.tight_layout()
    plt.savefig(chart_dir / 'weather_usage.png', dpi=150)
    plt.close()


def create_excel(df: pd.DataFrame, insights: dict, metrics: dict) -> Path:
    path = OUTPUT_DIR / 'water_usage_insight_premium.xlsx'
    monthly = df.assign(month=df['date'].dt.to_period('M').astype(str)).groupby('month', as_index=False).agg(
        average_usage=('usage', 'mean'),
        anomaly_count=('model_anomaly', 'sum'),
        leak_period_readings=('leak_report', 'sum'),
    )
    locality = df.groupby('locality', as_index=False).agg(
        average_usage=('usage', 'mean'),
        readings=('usage', 'size'),
        anomalies=('model_anomaly', 'sum'),
        leak_period_readings=('leak_report', 'sum'),
    ).sort_values('average_usage', ascending=False)
    connection_type = df.groupby('connection_type', as_index=False).agg(
        average_usage=('usage', 'mean'),
        readings=('usage', 'size'),
        anomalies=('model_anomaly', 'sum'),
    ).sort_values('average_usage', ascending=False)
    anomaly_locality = df.groupby('locality', as_index=False).agg(
        model_anomalies=('model_anomaly', 'sum'),
        leak_period_readings=('leak_report', 'sum'),
    ).sort_values('model_anomalies', ascending=False)
    repair_summary = df[df['leak_report'] == 1].groupby('connection_id', as_index=False).agg(
        leak_period_readings=('leak_report', 'sum'),
        average_usage=('usage', 'mean'),
        maximum_usage=('usage', 'max'),
        model_anomalies=('model_anomaly', 'sum'),
    ).sort_values(['model_anomalies', 'maximum_usage'], ascending=False)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        pd.DataFrame([insights]).to_excel(writer, sheet_name='Executive Summary', index=False)
        pd.DataFrame([metrics]).drop(columns=['feature_columns']).to_excel(writer, sheet_name='Model Metrics', index=False)
        df.sort_values('anomaly_score', ascending=False).head(250)[[
            'date', 'connection_id', 'locality', 'connection_type', 'usage',
            'anomaly_score', 'model_anomaly', 'leak_report', 'meter_status',
        ]].to_excel(writer, sheet_name='Top Anomalies', index=False)
        df.groupby(['locality', 'connection_type'], as_index=False).agg(
            average_usage=('usage', 'mean'),
            readings=('usage', 'size'),
            anomalies=('model_anomaly', 'sum'),
        ).to_excel(writer, sheet_name='Segment Summary', index=False)
        monthly.to_excel(writer, sheet_name='Monthly Trend', index=False)
        locality.to_excel(writer, sheet_name='Locality Summary', index=False)
        connection_type.to_excel(writer, sheet_name='Connection Summary', index=False)
        anomaly_locality.to_excel(writer, sheet_name='Anomaly Summary', index=False)
        repair_summary.to_excel(writer, sheet_name='Repair Analysis', index=False)
    workbook = __import__('openpyxl').load_workbook(path)
    dashboard = workbook.create_sheet('Dashboard', 0)
    dashboard.sheet_view.showGridLines = False
    dashboard.freeze_panes = None
    dashboard.sheet_properties.tabColor = '167C80'
    dark = PatternFill('solid', fgColor='17324D')
    teal = PatternFill('solid', fgColor='167C80')
    pale = PatternFill('solid', fgColor='E8F3F2')
    orange = PatternFill('solid', fgColor='F4B183')
    white_font = Font(color='FFFFFF', bold=True)
    title_font = Font(color='FFFFFF', bold=True, size=20)
    card_font = Font(color='17324D', bold=True, size=16)
    thin_border = Border(bottom=Side(style='thin', color='D9E2F3'))

    dashboard.merge_cells('B2:S2')
    dashboard['B2'] = 'WATER USAGE INSIGHT | EXCEL DASHBOARD'
    dashboard['B2'].fill = dark
    dashboard['B2'].font = title_font
    dashboard['B2'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard.row_dimensions[2].height = 32
    dashboard.merge_cells('B3:S3')
    dashboard['B3'] = f"Analysis period: {insights['date_start']} to {insights['date_end']}  |  Use the tabs below for detail"
    dashboard['B3'].font = Font(italic=True, color='5B6573')

    cards = [
        ('B5:D5', 'B6:D7', 'AVERAGE USAGE', f"{insights['mean_usage']:,.1f}"),
        ('F5:H5', 'F6:H7', 'READINGS ANALYZED', f"{insights['row_count']:,}"),
        ('J5:L5', 'J6:L7', 'MODEL ANOMALIES', f"{insights['model_anomalies']:,}"),
        ('N5:P5', 'N6:P7', 'LEAK-PERIOD READINGS', f"{insights['reported_leak_readings']:,}"),
    ]
    for label_range, value_range, label, value in cards:
        dashboard.merge_cells(label_range)
        dashboard.merge_cells(value_range)
        label_cell = dashboard[label_range.split(':')[0]]
        value_cell = dashboard[value_range.split(':')[0]]
        label_cell.value = label
        label_cell.fill = teal
        label_cell.font = white_font
        label_cell.alignment = Alignment(horizontal='center')
        value_cell.value = value
        value_cell.fill = pale if label != 'MODEL ANOMALIES' else orange
        value_cell.font = card_font
        value_cell.alignment = Alignment(horizontal='center', vertical='center')
    dashboard.row_dimensions[6].height = 23
    dashboard.row_dimensions[7].height = 23

    dashboard.merge_cells('B9:J9')
    dashboard['B9'] = 'MONTHLY USAGE TREND'
    dashboard['B9'].font = Font(bold=True, color='17324D', size=13)
    dashboard.merge_cells('K9:S9')
    dashboard['K9'] = 'AVERAGE USAGE BY LOCALITY'
    dashboard['K9'].font = Font(bold=True, color='17324D', size=13)
    dashboard['B9'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard['K9'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard.row_dimensions[9].height = 28
    dashboard.merge_cells('B27:J27')
    dashboard['B27'] = 'AVERAGE USAGE BY CONNECTION TYPE'
    dashboard['B27'].font = Font(bold=True, color='17324D', size=13)
    dashboard['B27'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard.merge_cells('K27:S27')
    dashboard['K27'] = 'MODEL ANOMALIES BY LOCALITY'
    dashboard['K27'].font = Font(bold=True, color='17324D', size=13)
    dashboard['K27'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard.merge_cells('B45:S45')
    dashboard['B45'] = 'MONTHLY ANOMALY TREND'
    dashboard['B45'].font = Font(bold=True, color='17324D', size=13)
    dashboard['B45'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard.merge_cells('B63:S63')
    dashboard['B63'] = 'INVESTIGATION GUIDANCE'
    dashboard['B63'].font = Font(bold=True, color='17324D', size=13)
    dashboard['B63'].alignment = Alignment(horizontal='center', vertical='center')
    guidance = [
        f"1. Start with {insights['top_locality']}, the highest-use locality ({insights['top_locality_mean_usage']:.1f} average).",
        '2. Review Top Anomalies with leak reports and meter status before dispatching field work.',
        f"3. Commercial connections average {insights['connection_type_means'].get('Commercial', 0):.1f}; compare them separately from residential connections.",
        '4. Treat model flags as investigation leads, not confirmed leaks; current precision is low against generated labels.',
    ]
    for row, message in enumerate(guidance, start=64):
        dashboard.merge_cells(start_row=row, start_column=2, end_row=row, end_column=19)
        cell = dashboard.cell(row=row, column=2, value=message)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.border = thin_border
    dashboard.merge_cells('B69:S69')
    dashboard['B69'] = 'Workbook guide: Dashboard = overview | Summary tabs = comparisons | Top Anomalies = field queue'
    dashboard['B69'].font = Font(italic=True, color='5B6573')

    monthly_sheet = workbook['Monthly Trend']
    line = LineChart()
    line.title = None
    line.y_axis.title = None
    line.x_axis.title = None
    line.height = 7.5
    line.width = 15.5
    line.style = 13
    line.legend = None
    line.add_data(Reference(monthly_sheet, min_col=2, max_col=2, min_row=1, max_row=monthly_sheet.max_row), titles_from_data=True)
    line.set_categories(Reference(monthly_sheet, min_col=1, min_row=2, max_row=monthly_sheet.max_row))
    dashboard.add_chart(line, 'B10')

    locality_sheet = workbook['Locality Summary']
    bar = BarChart()
    bar.type = 'bar'
    bar.style = 10
    bar.title = None
    bar.x_axis.title = None
    bar.y_axis.title = None
    bar.height = 7.5
    bar.width = 15.5
    bar.legend = None
    bar.gapWidth = 35
    bar.add_data(Reference(locality_sheet, min_col=2, max_col=2, min_row=1, max_row=locality_sheet.max_row), titles_from_data=True)
    bar.set_categories(Reference(locality_sheet, min_col=1, min_row=2, max_row=locality_sheet.max_row))
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.dataLabels.showSerName = False
    bar.dataLabels.showCatName = False
    bar.dataLabels.showLegendKey = False
    bar.dataLabels.showPercent = False
    bar.dataLabels.showLeaderLines = False
    bar.dataLabels.dLblPos = 'outEnd'
    dashboard.add_chart(bar, 'K10')

    connection_sheet = workbook['Connection Summary']
    connection_chart = BarChart()
    connection_chart.type = 'col'
    connection_chart.style = 10
    connection_chart.title = None
    connection_chart.x_axis.title = None
    connection_chart.y_axis.title = None
    connection_chart.height = 7.5
    connection_chart.width = 15.5
    connection_chart.legend = None
    connection_chart.gapWidth = 45
    connection_chart.add_data(Reference(connection_sheet, min_col=2, max_col=2, min_row=1, max_row=connection_sheet.max_row), titles_from_data=True)
    connection_chart.set_categories(Reference(connection_sheet, min_col=1, min_row=2, max_row=connection_sheet.max_row))
    connection_chart.dataLabels = DataLabelList()
    connection_chart.dataLabels.showVal = True
    connection_chart.dataLabels.showSerName = False
    connection_chart.dataLabels.showCatName = False
    connection_chart.dataLabels.showLegendKey = False
    connection_chart.dataLabels.dLblPos = 'outEnd'
    dashboard.add_chart(connection_chart, 'B28')

    anomaly_sheet = workbook['Anomaly Summary']
    anomaly_chart = BarChart()
    anomaly_chart.type = 'bar'
    anomaly_chart.style = 10
    anomaly_chart.title = None
    anomaly_chart.x_axis.title = None
    anomaly_chart.y_axis.title = None
    anomaly_chart.height = 7.5
    anomaly_chart.width = 15.5
    anomaly_chart.legend = None
    anomaly_chart.gapWidth = 35
    anomaly_chart.add_data(Reference(anomaly_sheet, min_col=2, max_col=2, min_row=1, max_row=anomaly_sheet.max_row), titles_from_data=True)
    anomaly_chart.set_categories(Reference(anomaly_sheet, min_col=1, min_row=2, max_row=anomaly_sheet.max_row))
    anomaly_chart.dataLabels = DataLabelList()
    anomaly_chart.dataLabels.showVal = True
    anomaly_chart.dataLabels.showSerName = False
    anomaly_chart.dataLabels.showCatName = False
    anomaly_chart.dataLabels.showLegendKey = False
    anomaly_chart.dataLabels.dLblPos = 'outEnd'
    dashboard.add_chart(anomaly_chart, 'K28')

    anomaly_trend = LineChart()
    anomaly_trend.title = None
    anomaly_trend.y_axis.title = None
    anomaly_trend.x_axis.title = None
    anomaly_trend.height = 7.5
    anomaly_trend.width = 15.5
    anomaly_trend.style = 13
    anomaly_trend.legend = None
    anomaly_trend.add_data(Reference(monthly_sheet, min_col=3, max_col=3, min_row=1, max_row=monthly_sheet.max_row), titles_from_data=True)
    anomaly_trend.set_categories(Reference(monthly_sheet, min_col=1, min_row=2, max_row=monthly_sheet.max_row))
    dashboard.add_chart(anomaly_trend, 'B46')

    for row in range(1, 70):
        for column in range(1, 20):
            dashboard.cell(row=row, column=column).alignment = Alignment(
                horizontal='center',
                vertical='center',
                wrap_text=True,
            )
    dashboard['B2'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard['B3'].alignment = Alignment(horizontal='center', vertical='center')
    dashboard['B33'].alignment = Alignment(horizontal='center', vertical='center')

    for sheet in workbook.worksheets:
        sheet.freeze_panes = None if sheet.title == 'Dashboard' else 'A2'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = teal
        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            values = [str(cell.value) for cell in column_cells[:50] if cell.value is not None]
            width = min(max([len(value) for value in values] + [10]) + 2, 28)
            sheet.column_dimensions[column_letter].width = width
    for column in ['B', 'C', 'D', 'F', 'G', 'H', 'J', 'K', 'L', 'N', 'O', 'P']:
        dashboard.column_dimensions[column].width = 16
    dashboard.column_dimensions['A'].width = 3
    dashboard.column_dimensions['D'].width = 3
    dashboard.column_dimensions['G'].width = 3
    dashboard.column_dimensions['J'].width = 3
    dashboard.column_dimensions['M'].width = 3
    dashboard.column_dimensions['Q'].width = 3
    dashboard.column_dimensions['R'].width = 3
    dashboard.column_dimensions['S'].width = 3
    dashboard.sheet_view.zoomScale = 85
    dashboard.print_area = 'B2:S69'
    dashboard.freeze_panes = None
    workbook.active = 0
    workbook.save(path)
    return path


def _add_title(slide, title: str, subtitle: str = '') -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.8))
    frame = box.text_frame
    frame.text = title
    frame.paragraphs[0].font.size = Pt(26)
    frame.paragraphs[0].font.bold = True
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.62), Inches(1.1), Inches(11), Inches(0.45))
        sub.text_frame.text = subtitle
        sub.text_frame.paragraphs[0].font.size = Pt(13)


def create_powerpoint(insights: dict, metrics: dict) -> Path:
    path = OUTPUT_DIR / 'water_usage_insight.pptx'
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_title(slide, 'Water Usage Insight', 'Usage monitoring, anomaly detection, and conservation priorities')
    body = slide.shapes.add_textbox(Inches(0.8), Inches(2), Inches(11.5), Inches(3.8))
    body.text_frame.text = f"Analysis period: {insights['date_start']} to {insights['date_end']}\nRows analyzed: {insights['row_count']:,}\nAverage daily reading: {insights['mean_usage']:.1f}\nModel-flagged anomalies: {insights['model_anomalies']:,}"
    body.text_frame.paragraphs[0].font.size = Pt(22)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_title(slide, 'What the data says')
    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(4.8))
    body.text_frame.text = f"Highest average locality: {insights['top_locality']} ({insights['top_locality_mean_usage']:.1f})\nReported leak-period readings: {insights['reported_leak_readings']:,}\nTop anomaly connection: {insights['top_anomaly_connection'] or 'None'}\nResidential average: {insights['connection_type_means'].get('Residential', 0):.1f}\nCommercial average: {insights['connection_type_means'].get('Commercial', 0):.1f}"
    body.text_frame.paragraphs[0].font.size = Pt(20)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_title(slide, 'Model validation and action')
    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(4.8))
    body.text_frame.text = f"Isolation Forest precision: {metrics['precision']:.3f}\nRecall against generated investigation labels: {metrics['recall']:.3f}\nF1 score: {metrics['f1']:.3f}\n\nPrioritize field checks for repeated anomalies, validate meters with extreme readings, and monitor usage after maintenance events."
    body.text_frame.paragraphs[0].font.size = Pt(20)
    presentation.save(path)
    return path


def generate_reports(df: pd.DataFrame, metrics: dict) -> dict:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    insights = build_insights(df)
    create_charts(df)
    excel_path = create_excel(df, insights, metrics)
    pptx_path = create_powerpoint(insights, metrics)
    with open(OUTPUT_DIR / 'insights.json', 'w', encoding='utf-8') as handle:
        json.dump(insights, handle, indent=2)
    return {'insights': insights, 'excel': str(excel_path), 'powerpoint': str(pptx_path)}
